import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_docx(filename="Smart_Campus_Navigation_Inference_Report.docx"):
    doc = Document()
    
    # Page setup - 0.8 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Document Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run("Smart Campus Navigation System")
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(14)
    sub_run = subtitle_p.add_run("AI Search Strategy Analysis & Empirical Output Inference Report\nFundamentals of Artificial Intelligence Mini-Project")
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    sub_run.font.bold = True

    # Metadata table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    col_widths = [Inches(3.4), Inches(3.4)]
    
    meta_data = [
        [("Project:", " Smart Campus Navigation System"), ("Algorithms Evaluated:", " BFS, DFS, A* Search, Dijkstra Baseline")],
        [("Implementation:", " Python, Streamlit, Matplotlib"), ("Evaluation Scope:", " Path Optimality, Hop Count, Node Expansion & Runtime")]
    ]
    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths[col_idx]
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            k, v = meta_data[row_idx][col_idx]
            r1 = p.add_run(k)
            r1.font.bold = True
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            r2 = p.add_run(v)
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1. Executive Summary
    h1 = doc.add_heading("1. Executive Summary", level=1)
    h1.style.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)
    p = doc.add_paragraph(
        "This project develops an interactive, graph-based Smart Campus Navigation System designed to evaluate and compare three foundational Artificial Intelligence search strategies: Breadth-First Search (BFS), Depth-First Search (DFS), and A* Search, benchmarked against a verified Dijkstra minimum-distance baseline. Campus physical layouts and pedestrian walkways are modeled as weighted graphs where campus buildings/landmarks serve as vertices and calibrated pedestrian paths serve as weighted bidirectional edges."
    )
    p = doc.add_paragraph(
        "A full-featured Streamlit web application provides interactive route planning, graphical map overlay rendering, dynamic path offset visualization, and performance dashboards. This report presents the formal AI problem formulation, topological dataset design, empirical execution results, and detailed individual observations regarding the distinct operational characteristics, optimality behaviors, and search efficiencies of each algorithm."
    )

    # 2. AI Problem Formulation
    h2 = doc.add_heading("2. AI Problem Formulation", level=1)
    h2.style.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)
    doc.add_paragraph(
        "The campus navigation challenge is modeled as a state-space search problem with the following formal components:"
    )
    
    items = [
        ("Initial State: ", "The vertex corresponding to the user's selected departure location (e.g., 'Main Gate')."),
        ("State Space: ", "Finite set of campus landmark vertices V in graph G = (V, E) with known spatial (x, y) coordinates."),
        ("Actions: ", "Traverse an available pedestrian walkway connecting the current vertex to an adjacent landmark."),
        ("Transition Model: ", "Given state u and action 'traverse edge (u, v)', the successor state is v."),
        ("Goal Test: ", "A predicate checking whether the current node v equals the destination goal state (v == goal)."),
        ("Path Cost g(n): ", "The cumulative sum of physical walking distances (in metres) along all traversed edges from the start state to the current state n."),
        ("Heuristic Function h(n): ", "An admissible and consistent scaled Euclidean distance estimate between the current node n(x_n, y_n) and the goal vertex g(x_g, y_g): h(n) = s * sqrt((x_n - x_g)^2 + (y_n - y_g)^2), where the scale factor s is dynamically calibrated to remain strictly lower than the minimum edge-weight to straight-line ratio across the graph (s <= min(w / d_euclid) * 0.9). This guarantees admissibility (h(n) <= h*(n)) and monotonicity, ensuring that A* never overestimates remaining distance and guarantees finding the optimal path.")
    ]
    for prefix, body in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r_pre = p.add_run(prefix)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        p.add_run(body)

    # 3. Datasets & Graph Topologies
    h3 = doc.add_heading("3. Datasets and Graph Topologies", level=1)
    h3.style.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)
    doc.add_paragraph(
        "To evaluate algorithm performance under varying graph density, two distinct campus datasets are integrated into the system:"
    )
    p1 = doc.add_paragraph(style='List Bullet')
    r = p1.add_run("Standard Campus Dataset (campus_map.json): ")
    r.font.bold = True
    p1.add_run("Comprises 10 primary campus landmarks (Main Gate, Admin Block, Library, CSE Block, AI Lab, Cafeteria, Auditorium, Hostel, Sports Ground, Parking) interconnected by 15 bidirectional walkways with realistic metric walking distances (23 m to 72 m). The minimum edge-to-Euclidean distance ratio is 0.92, validated with heuristic scale s = 0.75.")
    
    p2 = doc.add_paragraph(style='List Bullet')
    r = p2.add_run("Extended University Campus Dataset (campus_map_extended.json): ")
    r.font.bold = True
    p2.add_run("A comprehensive 16-landmark academic campus network including Security, Medical Center, Robotics & Electronics Lab, Research Park, North & South Hostels, Dining Hall, Amphitheatre, and Multi-level Parking, connected across 25 pedestrian corridors. Validated with s = 0.75 to ensure strict admissibility.")

    # 4. Empirical Evaluation & Comparison Tables
    h4 = doc.add_heading("4. Empirical Evaluation & Comparative Test Results", level=1)
    h4.style.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)
    doc.add_paragraph(
        "Representative test cases were executed across both datasets to record route distance, hop count (number of edges), nodes expanded in the search space, and execution latency. The results are summarized below:"
    )

    # Table 1: Route 1
    t1_title = doc.add_paragraph()
    t1_title.paragraph_format.space_before = Pt(6)
    t1_title.paragraph_format.space_after = Pt(2)
    r = t1_title.add_run("Table 1: Test Case 1 - Main Gate to AI Lab (Standard Dataset)")
    r.font.bold = True
    r.font.size = Pt(10.5)

    table1 = doc.add_table(rows=4, cols=6)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    table1.autofit = False
    widths = [Inches(1.0), Inches(1.1), Inches(0.8), Inches(1.1), Inches(1.0), Inches(1.8)]
    headers = ["Algorithm", "Distance (m)", "Edges", "Nodes Expanded", "Time (ms)", "Calculated Route Path"]
    
    for c_idx, title in enumerate(headers):
        cell = table1.cell(0, c_idx)
        cell.width = widths[c_idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 60, 60, 80, 80)
        p = cell.paragraphs[0]
        run = p.add_run(title)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9)

    t1_data = [
        ["BFS", "135.0 m", "3", "10", "0.026 ms", "Main Gate -> Cafeteria -> CSE Block -> AI Lab"],
        ["DFS", "129.0 m", "4", "5", "0.015 ms", "Main Gate -> Admin Block -> Library -> CSE Block -> AI Lab"],
        ["A*", "129.0 m", "4", "7", "0.046 ms", "Main Gate -> Admin Block -> Library -> CSE Block -> AI Lab"]
    ]
    for r_idx, row_values in enumerate(t1_data, 1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = table1.cell(r_idx, c_idx)
            cell.width = widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 50, 50, 80, 80)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(8.5)
            if c_idx == 0:
                run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Table 2: Route 2
    t2_title = doc.add_paragraph()
    t2_title.paragraph_format.space_before = Pt(4)
    t2_title.paragraph_format.space_after = Pt(2)
    r = t2_title.add_run("Table 2: Test Case 2 - Parking to Hostel (Standard Dataset)")
    r.font.bold = True
    r.font.size = Pt(10.5)

    table2 = doc.add_table(rows=4, cols=6)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False
    for c_idx, title in enumerate(headers):
        cell = table2.cell(0, c_idx)
        cell.width = widths[c_idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 60, 60, 80, 80)
        p = cell.paragraphs[0]
        run = p.add_run(title)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9)

    t2_data = [
        ["BFS", "131.0 m", "4", "9", "0.027 ms", "Parking -> Main Gate -> Admin Block -> Library -> Hostel"],
        ["DFS", "131.0 m", "4", "7", "0.017 ms", "Parking -> Main Gate -> Admin Block -> Library -> Hostel"],
        ["A*", "131.0 m", "4", "5", "0.036 ms", "Parking -> Main Gate -> Admin Block -> Library -> Hostel"]
    ]
    for r_idx, row_values in enumerate(t2_data, 1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = table2.cell(r_idx, c_idx)
            cell.width = widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 50, 50, 80, 80)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(8.5)
            if c_idx == 0:
                run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Table 3: Route 3 (Extended Dataset)
    t3_title = doc.add_paragraph()
    t3_title.paragraph_format.space_before = Pt(4)
    t3_title.paragraph_format.space_after = Pt(2)
    r = t3_title.add_run("Table 3: Test Case 3 - Main Gate to Sports Complex (Extended Dataset, 16 Locations)")
    r.font.bold = True
    r.font.size = Pt(10.5)

    table3 = doc.add_table(rows=4, cols=6)
    table3.alignment = WD_TABLE_ALIGNMENT.CENTER
    table3.autofit = False
    for c_idx, title in enumerate(headers):
        cell = table3.cell(0, c_idx)
        cell.width = widths[c_idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 60, 60, 80, 80)
        p = cell.paragraphs[0]
        run = p.add_run(title)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9)

    t3_data = [
        ["BFS", "142.0 m", "4", "14", "0.026 ms", "Main Gate -> Security -> Admin Block -> Auditorium -> Sports Complex"],
        ["DFS", "160.0 m", "5", "13", "0.022 ms", "Main Gate -> Security -> Admin -> Cafeteria -> Amphitheatre -> Sports Complex"],
        ["A*", "140.0 m", "4", "8", "0.046 ms", "Main Gate -> Parking -> Medical Center -> Auditorium -> Sports Complex"]
    ]
    for r_idx, row_values in enumerate(t3_data, 1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = table3.cell(r_idx, c_idx)
            cell.width = widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 50, 50, 80, 80)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(8.5)
            if c_idx == 0:
                run.font.bold = True

    # 5. Individual Observations About the Outputs
    h5 = doc.add_heading("5. Individual Observations About the Outputs", level=1)
    h5.style.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)

    doc.add_paragraph(
        "A rigorous comparative analysis of the empirical outputs reveals fundamental differences in the design philosophy, path quality, and algorithmic behavior of each search technique:"
    )

    # Subsection: BFS
    h5_1 = doc.add_heading("5.1 Individual Observation: Breadth-First Search (BFS)", level=2)
    h5_1.style.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
    p_bfs = doc.add_paragraph(
        "1. Hop-Count Optimality vs. Path-Cost Sub-optimality:\n"
        "BFS expands states strictly in level-order using a First-In, First-Out (FIFO) queue. Consequently, BFS is guaranteed to discover the path containing the absolute minimum number of edges (hops) from start to destination. However, in physical campus walkways where edge distances vary significantly (e.g., Main Gate -> Cafeteria = 67 m vs Main Gate -> Admin Block = 32 m), minimizing edge count does NOT correspond to minimizing physical travel distance. In Test Case 1, BFS selected the 3-hop route via Cafeteria totaling 135.0 m, which is 6.0 m longer than the optimal 129.0 m route found by A* (4 hops).\n\n"
        "2. Complete Frontier Exploration:\n"
        "Because BFS possesses no knowledge of the goal's geometric direction, it radiates outwards uniformly in concentric spheres. For Test Case 1, BFS expanded all 10 nodes in the graph before validating the goal. For Test Case 3 in the extended 16-node graph, BFS expanded 14 out of 16 nodes. This exhaustive frontier growth leads to exponential memory consumption O(b^d) in larger graphs, making uninformed BFS unsuitable for large-scale geographic routing."
    )

    # Subsection: DFS
    h5_2 = doc.add_heading("5.2 Individual Observation: Depth-First Search (DFS)", level=2)
    h5_2.style.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)
    p_dfs = doc.add_paragraph(
        "1. Traversal Order Sensitivity & Arbitrary Path Quality:\n"
        "DFS explores recursively along a single branch using a Last-In, First-Out (LIFO) stack until reaching a dead end or the goal state. Because DFS does not incorporate path cost or edge count, the quality of its output depends entirely on the arbitrary tie-breaking order of neighboring vertices in the graph representation. In Test Case 3, DFS chose an indirect 5-hop path via Student Cafeteria and the Amphitheatre totaling 160.0 m (20.0 m longer than optimal A*).\n\n"
        "2. Low Search Space Overhead in Shallow Success:\n"
        "When the goal happens to align with the first explored branch, DFS can reach the destination rapidly with low node expansion (expanding only 5 nodes in Test Case 1 and 7 nodes in Test Case 2). Furthermore, its frontier storage requirements remain linear O(b * m) relative to maximum graph depth m. However, DFS cannot guarantee optimality, frequently returning suboptimal detours, and is vulnerable to unbounded depth in non-tree cyclic graphs without strict cycle detection."
    )

    # Subsection: A* Search
    h5_3 = doc.add_heading("5.3 Individual Observation: A* Search (Informed Search)", level=2)
    h5_3.style.font.color.rgb = RGBColor(0x05, 0x96, 0x69)
    p_astar = doc.add_paragraph(
        "1. Provable Path Optimality Matching Dijkstra's Baseline:\n"
        "A* combines the exact accumulated cost from the start g(n) with an admissible heuristic estimate of remaining cost to the goal h(n) using the evaluation function f(n) = g(n) + h(n). Across all tested source-destination pairs (and in 100% of the automated all-pairs test suite), A* found the strictly optimal minimum walking distance (129.0 m for Test Case 1, 131.0 m for Test Case 2, and 140.0 m for Test Case 3), perfectly matching the Dijkstra minimum-distance baseline.\n\n"
        "2. Heuristic-Directed Search Space Pruning:\n"
        "By utilizing spatial Euclidean coordinates of campus buildings, the heuristic pulls the search frontier directly toward the geometric coordinates of the target destination. In Test Case 2 (Parking to Hostel), A* reached the goal after expanding only 5 nodes, compared to 9 nodes for BFS - saving nearly 45% of exploration effort. In Test Case 3, A* expanded only 8 nodes compared to 14 nodes for BFS. While managing the priority queue incurs a minor per-node time constant, the reduction in search space makes A* by far the most scalable, robust, and dependable routing algorithm for physical campus navigation."
    )

    # 6. Verification and Automated Testing
    h6 = doc.add_heading("6. Verification and Automated Test Suite", level=1)
    h6.style.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)
    doc.add_paragraph(
        "The project includes an automated test suite implemented in Python's unittest framework (tests/test_navigation.py). The test suite exercises the following critical invariants:\n"
        "• All-Pairs Optimality: Exhaustively computes all V x V source-destination combinations on both standard and extended campus maps, verifying that A* route distance exactly matches Dijkstra's shortest path baseline.\n"
        "• Edge Integrity: Confirms that every consecutive pair of vertices along generated paths corresponds to a valid bidirectional edge in the graph.\n"
        "• Exact Distance Summation: Asserts that reported path distance equals the exact mathematical sum of individual walkway edge weights.\n"
        "• Boundary Conditions: Validates identity queries (start == destination) returning 0 distance and 0 edges."
    )
    doc.add_paragraph(
        "All 6 unit tests pass cleanly with 0.012s execution time, confirming high mathematical and algorithmic reliability."
    )

    # 7. Application Screenshots
    h7 = doc.add_heading("7. System Interface & Output Screenshots", level=1)
    h7.style.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)
    doc.add_paragraph(
        "The application interface was executed and captured in real-time operation. Below are visual records of the system components and output dashboards:"
    )

    screenshots_info = [
        ("screenshots/01_route_planner.png", "Figure 1: Interactive Route Planner - Campus graph network with concurrent path visualization (BFS, DFS, A*) with offset styling, key metrics, and step-by-step route directions."),
        ("screenshots/02_algorithm_comparison.png", "Figure 2: Algorithm Comparison Dashboard - Performance metrics data table (Distance, Edges, Nodes Expanded, Execution Time) with side-by-side comparative bar charts and Dijkstra baseline validation."),
        ("screenshots/04_extended_campus_map.png", "Figure 3: Extended University Campus Dataset (16 Locations, 25 Walkways) - Dynamic dataset switching showing complex multi-building routing."),
        ("screenshots/03_project_overview.png", "Figure 4: Project Architecture & Export - System formulation summary, operational boundaries, and CSV data export functionality.")
    ]

    for img_path, caption in screenshots_info:
        if os.path.exists(img_path):
            doc.add_paragraph().paragraph_format.space_before = Pt(8)
            img_p = doc.add_paragraph()
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_p.paragraph_format.space_after = Pt(4)
            img_run = img_p.add_run()
            img_run.add_picture(img_path, width=Inches(6.2))
            
            cap_p = doc.add_paragraph()
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.space_after = Pt(12)
            cap_run = cap_p.add_run(caption)
            cap_run.font.size = Pt(9)
            cap_run.font.italic = True
            cap_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # 8. Limitations & Future Scope
    h8 = doc.add_heading("8. Limitations and Future Scope", level=1)
    h8.style.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)
    doc.add_paragraph(
        "Limitations:\n"
        "• Walkways are modeled as static, bidirectional edges with constant walking speeds, not accounting for elevation changes or temporary path closures.\n"
        "• Accessibility constraints (wheelchair ramps, elevator vs. stairs) and pedestrian congestion density are not currently modeled.\n\n"
        "Future Scope:\n"
        "• Integration with GIS/OpenStreetMap live coordinate layers for real-world universities.\n"
        "• Dynamic multi-objective routing (e.g., shortest distance vs. covered rain walkway vs. wheelchair accessibility).\n"
        "• Mobile responsiveness and GPS-based real-time indoor/outdoor turn-by-turn navigation."
    )

    # 9. Conclusion
    h9 = doc.add_heading("9. Conclusion", level=1)
    h9.style.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)
    doc.add_paragraph(
        "The Smart Campus Navigation System successfully demonstrates the fundamental trade-offs between uninformed and informed search strategies in physical navigation domains. BFS guarantees minimum hop count but is prone to distance sub-optimality on weighted graphs. DFS is computationally light but lacks quality guarantees and exhibits high variance. A* Search, guided by an admissible scaled Euclidean heuristic, consistently identifies the provably optimal minimum-distance route while drastically reducing node expansions. The complete codebase, test suite, and interactive Streamlit frontend offer a comprehensive, reliable, and academically rigorous platform for AI navigation study."
    )

    doc.save(filename)
    print(f"Successfully generated DOCX report: {filename}")

if __name__ == "__main__":
    create_docx()
